import io

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from cccc_desensitization import embedded_images, file_handlers, main


def sample():
    image = io.BytesIO()
    Image.new("RGB", (40, 30), "white").save(image, "PNG")
    document = Document()
    run = document.add_paragraph().add_run("secret")
    run.bold = True
    run.add_picture(io.BytesIO(image.getvalue()))
    document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].add_run().add_picture(io.BytesIO(image.getvalue()))
    document.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(image.getvalue()))
    floating = document.sections[0].footer.paragraphs[0].add_run().add_picture(io.BytesIO(image.getvalue()))
    floating._inline.tag = qn("wp:anchor")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue(), image.getvalue()


def test_shared_images_and_text_nodes_survive(tmp_path):
    data, original = sample()
    calls = []
    def ocr(blob, folder):
        calls.append(blob)
        return 40, 30, [{"id": "box", "page": 1, "x": 0, "y": 0, "width": .5, "height": .5, "selected": True}], [], []
    images, warnings = embedded_images.analyze_images(data, tmp_path, ocr)
    assert not warnings and calls == [original]
    assert len(images) == 1 and len(images[0]["locations"]) == 3
    entity = {"start": 0, "end": 6, "text": "secret", "type": "CUSTOM", "selected": True}
    text_masked = file_handlers.mask_file(data, "sample.docx", "secret", [entity], {"CUSTOM": {"text_action": "replace", "replacement": "hidden"}})
    result, warnings = embedded_images.mask_images(text_masked, images, {"DEFAULT": {"image_action": "fill", "color": "#000000"}}, file_handlers.mask_file)
    document = Document(io.BytesIO(result))
    assert document.paragraphs[0].text == "hidden"
    assert document.paragraphs[0].runs[0].bold
    assert len(list(document.paragraphs[0]._p.iter(qn("w:drawing")))) == 1
    assert len(embedded_images.resources(document)) == 1
    blob = embedded_images.resources(document)[0][1].blob
    assert Image.open(io.BytesIO(blob)).convert("RGB").getpixel((3, 3)) == (0, 0, 0)
    assert not warnings


def test_image_failure_keeps_original_without_sensitive_warning(tmp_path):
    data, original = sample()
    def failing(*args):
        raise RuntimeError("sensitive original text")
    images, warnings = embedded_images.analyze_images(data, tmp_path, failing)
    assert images[0]["status"] == "failed"
    assert "sensitive original text" not in str(warnings)
    result, _ = embedded_images.mask_images(data, images, {}, failing)
    assert embedded_images.resources(Document(io.BytesIO(result)))[0][1].blob == original


def test_review_rejects_foreign_image_and_coordinates():
    image = {"image_id": "a", "entities": [], "boxes": [], "status": "ready"}
    review = lambda submitted: embedded_images.review_images([image], submitted, main._review_entities, main._validate_boxes, main._synchronize_review_selection)
    with pytest.raises(ValueError):
        review([{**image, "image_id": "b"}])
    with pytest.raises(ValueError):
        review([])
    with pytest.raises(main.HTTPException):
        review([{**image, "boxes": [{"x": .9, "y": 0, "width": .2, "height": .2}]}])
    assert review([image])[0]["boxes"] == []


def test_adjusted_region_and_deselection(tmp_path):
    data, original = sample()
    box = {"id": "box", "x": 0, "y": 0, "width": .5, "height": .5, "selected": True}
    entity = {"id": "entity", "type": "PERSON", "text": "alice", "start": 0, "end": 5,
              "bbox": {"x": 0, "y": 0, "width": .5, "height": .5}, "box_ids": ["box"], "selected": True}
    images, _ = embedded_images.analyze_images(data, tmp_path, lambda *args: (40, 30, [box], [entity], []))
    submitted = [{**images[0], "boxes": [{**box, "x": .5}]}]
    reviewed = embedded_images.review_images(images, submitted, main._review_entities, main._validate_boxes, main._synchronize_review_selection)
    assert reviewed[0]["entities"][0]["bbox"]["x"] == .5
    result, _ = embedded_images.mask_images(data, reviewed, {"PERSON": {"image_action": "fill", "color": "#000000", "show_replacement": False}}, file_handlers.mask_file)
    blob = embedded_images.resources(Document(io.BytesIO(result)))[0][1].blob
    pixels = Image.open(io.BytesIO(blob)).convert("RGB")
    assert pixels.getpixel((3, 3)) == (255, 255, 255)
    assert pixels.getpixel((30, 3)) == (0, 0, 0)
    submitted[0]["entities"] = [{**entity, "selected": False}]
    reviewed = embedded_images.review_images(images, submitted, main._review_entities, main._validate_boxes, main._synchronize_review_selection)
    result, _ = embedded_images.mask_images(data, reviewed, {}, file_handlers.mask_file)
    assert embedded_images.resources(Document(io.BytesIO(result)))[0][1].blob == original


def test_external_images_are_never_downloaded(tmp_path):
    from docx.opc.constants import RELATIONSHIP_TYPE
    document = Document()
    document.part.relate_to("https://invalid.example/private.png", RELATIONSHIP_TYPE.IMAGE, is_external=True)
    output = io.BytesIO()
    document.save(output)
    images, warnings = embedded_images.analyze_images(output.getvalue(), tmp_path, lambda *args: pytest.fail("external OCR"))
    assert images[0]["status"] == "failed"
    assert warnings[0]["code"] == "DOCX_EXTERNAL_IMAGE"
    assert "private.png" not in str(warnings)


def test_mask_failure_preserves_resource(tmp_path):
    data, original = sample()
    images, _ = embedded_images.analyze_images(data, tmp_path, lambda *args: (40, 30, [{"selected": True}], [], []))
    def fail(*args): raise RuntimeError("private text")
    result, warnings = embedded_images.mask_images(data, images, {}, fail)
    assert warnings[0]["code"] == "DOCX_IMAGE_MASK_FAILED"
    assert "private text" not in str(warnings)
    assert embedded_images.resources(Document(io.BytesIO(result)))[0][1].blob == original


@pytest.mark.parametrize("score,selected", [(.049, False), (.05, True), (.051, True)])
def test_default_threshold(score, selected, monkeypatch):
    class Connection:
        def execute(self, *args): return self
        def fetchall(self): return []
        def close(self): pass
    monkeypatch.setattr(main, "conn", Connection)
    monkeypatch.setattr(main, "_rule_settings", lambda: {})
    monkeypatch.setattr(main.ner_service, "analyze", lambda text: [{"text": text, "type": "PERSON", "start": 0, "end": len(text), "score": score}])
    entities = main.analyze("alice")
    assert main.NER_SELECTION_THRESHOLD == 0.05
    assert entities[0]["selected"] is selected
    reviewed = main._review_entities(entities, [{**entities[0], "selected": True}])
    assert reviewed[0]["selected"] is True
